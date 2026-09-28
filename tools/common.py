# Shared helpers for the plan builders.
#   python tools/<script>.py            -> signed school copy on the Desktop (names from local_signatures.json)
#   python tools/<script>.py --public   -> public copy inside this repo (no signatures, credit footer on every page)
import json
import os
import sys

from openpyxl.styles import Alignment, Font

PUBLIC = "--public" in sys.argv
TOOLS = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(TOOLS)
DESKTOP = os.path.join(os.path.expanduser("~"), "Desktop")
CREDIT = "Prepared by Ogün Aydın · CC BY 4.0: free to use, share and adapt with credit · github.com/ogunaydin89/ingilizce-6-materyaller"


def out_path(folder, public_name, school_name):
    """Where the .xlsx goes; the PDF is written next to it with the same name."""
    if PUBLIC:
        d = os.path.join(REPO, folder)
        os.makedirs(d, exist_ok=True)
        return os.path.join(d, public_name)
    return os.path.join(DESKTOP, school_name)


def signatures():
    """Signature texts for the school copy. The file stays on this PC (.gitignore)."""
    with open(os.path.join(TOOLS, "local_signatures.json"), encoding="utf-8") as f:
        return json.load(f)


def finish(ws, row, left_cols, right_cols, date, size=9):
    """Signed copy: signature block at `row` (4 rows). Public copy: credit line + page footer.
    Returns the last used row."""
    if PUBLIC:
        ws.merge_cells(start_row=row, start_column=left_cols[0], end_row=row, end_column=right_cols[1])
        c = ws.cell(row, left_cols[0], CREDIT)
        c.font = Font(name="Calibri", size=size - 1, italic=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 14
        ws.oddFooter.center.text = CREDIT
        ws.oddFooter.center.size = 7
        return row
    sig = signatures()
    ws.merge_cells(start_row=row, start_column=left_cols[0], end_row=row + 3, end_column=left_cols[1])
    ws.merge_cells(start_row=row, start_column=right_cols[0], end_row=row + 3, end_column=right_cols[1])
    ws.cell(row, left_cols[0], f"{date}\n\n{sig['teacher_name']}\n{sig['teacher_title']}")
    ws.cell(row, right_cols[0], f"{sig['approval_word']}\n{date}\n\n{sig['principal_name']}\n{sig['principal_title']}")
    for col in (left_cols[0], right_cols[0]):
        ws.cell(row, col).font = Font(name="Calibri", size=size, bold=True)
        ws.cell(row, col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for k in range(4):
        ws.row_dimensions[row + k].height = 16
    return row + 3


def export_pdf(xlsx):
    """Print the workbook to PDF with Excel (same result as printing it)."""
    import win32com.client
    pdf = os.path.splitext(xlsx)[0] + ".pdf"
    xl = win32com.client.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    try:
        wb = xl.Workbooks.Open(xlsx)
        pages = wb.Worksheets(1).PageSetup.Pages.Count
        wb.ExportAsFixedFormat(0, pdf)
        wb.Close(False)
    finally:
        xl.Quit()
    print("pdf", pdf, "pages", pages)
    return pdf
