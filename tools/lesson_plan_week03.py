# Weekly lesson plan for 6/A–6/B English, week 3 (28.09–02.10.2026): Revision part 2.
import openpyxl

from common import PUBLIC, export_pdf, finish, out_path
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = out_path("Grade-6_Week-03_2026-09-28--2026-10-02_English_Revision-Part-2",
               "Grade-6_Week-03_English_Lesson-Plan.xlsx",
               "2026-2027_6AB_Ingilizce_Gunluk_Plan_Hafta3.xlsx")

thin = Side(style="thin", color="808080")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
GREY = PatternFill("solid", fgColor="D9D9D9")
F = Font(name="Calibri", size=9)
FB = Font(name="Calibri", size=9, bold=True)
WRAP = Alignment(wrap_text=True, vertical="center")
CWRAP = Alignment(wrap_text=True, vertical="center", horizontal="center")

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Lesson Plan"
for col, w in zip("ABCD", (20, 55, 55, 42)):
    ws.column_dimensions[col].width = w

r = 1


def title(text, size=11):
    global r
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    c = ws.cell(r, 1, text)
    c.font = Font(name="Calibri", size=size, bold=True)
    c.alignment = CWRAP
    ws.row_dimensions[r].height = 15
    r += 1


def row(label, text, height):
    global r
    ws.cell(r, 1, label).font = FB
    ws.cell(r, 1).fill = GREY
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    ws.cell(r, 2, text).font = F
    for col in range(1, 5):
        ws.cell(r, col).border = BOX
        ws.cell(r, col).alignment = WRAP
    ws.row_dimensions[r].height = height
    r += 1


title("SAİMBEYLİ ŞEHİT BAYRAM ALİ YARDIMCI YATILI BÖLGE ORTAOKULU · 2026-2027 ACADEMIC YEAR")
title("YEAR 6 ENGLISH LESSON PLAN · WEEK 3 (28 September – 2 October 2026)")
r += 1

row("SUBJECT / CLASS", "English · 6/A – 6/B · Level A2.2 · 3 lessons (40 min)", 14)
row("LESSON TIMES",
    "6/A: Thursday 01.10 – 7th lesson (14:40)  ·  Friday 02.10 – 1st–2nd lessons (08:30–10:05)\n"
    "6/B: Thursday 01.10 – 5th lesson (12:50)  ·  Friday 02.10 – 4th–5th lessons (11:20–13:30)", 26)
row("THEME", "REVISION PROGRAMME – Part 2: Year 5 (A2.1) Themes 5–8 (Neighbourhood & City · Life in the World · Life in Nature · Universe & Future)", 26)
row("LEARNING OUTCOMES",
    "ENG.5.5–5.8: V1 (target vocabulary), G1 (target structures), P1 (pronunciation), with L3/R3 (making meaning) and S4/W4 (short own sentences) in the activities. "
    "Pupils can recall and use: comparatives, possessive ’s, there is/are, present simple · can (permission), have got, countable/uncountable (how much/how many) · "
    "can (ability), must, where-questions, superlatives · be going to (plans).", 38)
row("MATERIALS", "Interactive board (EBA / digital coursebook – the printed Year 6 books have not been delivered yet), picture cards, a simple menu card, worksheet (10-question revision quiz)", 14)
r += 1

hdr = ["LESSON", "TEACHING-LEARNING PROCESS", "", "ASSESSMENT"]
for col, h in enumerate(hdr, 1):
    c = ws.cell(r, col, h)
    c.font, c.fill, c.border, c.alignment = FB, GREY, BOX, CWRAP
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
ws.row_dimensions[r].height = 14
r += 1

LESSONS = [
    ("LESSON 1 (Thursday)\nNeighbourhood & City",
     "• Warm-up (5′): new teacher – quick “Hello, my name is… I live in…” chain.\n"
     "• Places in the neighbourhood/city (square, museum, castle, sports centre, park…): picture–word match on the board.\n"
     "• Structures: There is/are…, comparatives: pupils compare Saimbeyli and Adana (“Adana is bigger than Saimbeyli.”), possessive ’s.\n"
     "• Pairs: 3 sentences about their village/town.",
     "Observation during the chain and pair work; exit ticket: 1 “there is/are” + 1 comparative sentence."),
    ("LESSON 2 (Friday)\nLife in the World",
     "• Food vocabulary (bread, soup, salad, beans, jam, lemon…) with picture cards; countable/uncountable sort on the board.\n"
     "• How much / how many…? · have got / has got.\n"
     "• Restaurant role-play with a simple menu card: “Can I have…? Can I pay in cash?” (can for permission).",
     "Role-play observed with a short checklist (vocabulary, can-question, pronunciation)."),
    ("LESSON 3 (Friday)\nLife in Nature · Universe & Future",
     "• Wild animals and habitats: “What can it do?” guessing game (can/can’t), “Where do lions live?”.\n"
     "• Comparative/superlative animal facts (“The whale is the biggest animal.”); must (“Birds must fly south in winter.”).\n"
     "• Holidays: be going to – each pupil writes 3 sentences “Next summer I am going to…”.\n"
     "• Last 10′: 10-question revision quiz covering Themes 1–8 (diagnostic, not graded) → basis for Theme 1 next week.",
     "Written sentences + 10-question quiz: note pupils who need support before Theme 1."),
]
for label, body, assess in LESSONS:
    ws.cell(r, 1, label).font = FB
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws.cell(r, 2, body).font = F
    ws.cell(r, 4, assess).font = F
    for col in range(1, 5):
        ws.cell(r, col).border = BOX
        ws.cell(r, col).alignment = WRAP if col > 1 else CWRAP
    ws.row_dimensions[r].height = 66
    r += 1
r += 1

row("DIFFERENTIATION", "Support: picture cards, sentence-frame strips, Turkish hints when needed. Extension: fast finishers write 2 extra sentences / their own menu.", 26)
row("NOTES",
    "In the annual plan, weeks 2–3 are the revision programme (the MoNE curriculum revises the Year 5 themes in 2 blocks at the start of the year). "
    "Block 1 (Themes 1–4) was covered in week 2; if not, use Themes 1–4 this week instead. 6/A’s Thursday lesson is the last period of the day: keep it game-based.", 38)

r += 1
last = finish(ws, r, (1, 2), (3, 4), "28.09.2026", size=10)

ps = ws.page_setup
ps.paperSize = ws.PAPERSIZE_A4
ps.orientation = "landscape"
ws.sheet_properties.pageSetUpPr.fitToPage = True
ps.fitToWidth = 1
ps.fitToHeight = 1
ws.page_margins.left = ws.page_margins.right = 0.4
ws.page_margins.top = ws.page_margins.bottom = 0.4
ws.print_options.horizontalCentered = True
ws.print_area = f"A1:D{last}"
wb.save(OUT)
print("saved", OUT)
export_pdf(OUT)
